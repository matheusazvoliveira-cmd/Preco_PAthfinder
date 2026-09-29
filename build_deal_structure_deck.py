"""Builds MRS_Pathfinder_Deal_Structure_v3_Bilingual.pptx (EN slides 1-2, PT slides 3-4).

Visual base = slide 1 of MRS_Pathfinder_GDT_Deal_Structure_Triangle_Bilingual_Revised.pptx
(now archived in Arquivo_Versoes_Antigas/, superseded by this v3 deck).
Run: PYTHONIOENCODING=utf-8 python build_deal_structure_deck.py
"""
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(__file__).parent / "MRS_Pathfinder_Deal_Structure_v3_Bilingual.pptx"

INK, MUTED, RED = "142C4B", "5C6874", "BA3434"
BLUE, BLUE_BG = "2366AC", "E7F0FA"
ORANGE, ORANGE_BG = "DE7E22", "FFF3E2"
GREEN, GREEN_BG = "2D874F", "E7F6EB"
TEAL, TEAL_BG = "008C96", "E7F8F7"
FONT, FONT_H = "Aptos", "Aptos Display"
W, H = 13.333, 7.5
M = 0.45

TXT = {
    "EN": dict(
        title="MRS PATHFINDER IRON ORE CAR — DEAL STRUCTURE",
        sub="Ownership, component flow, refurbishment and final sale",
        usa=("Wabtec Digital USA", ["PN supplier", "Sells components to Digital Brazil"]),
        sup=("Supermetal", ["Receives railcar + components", "Refurbishes and integrates",
                            "Returns finished Pathfinder Iron Ore Car", "Cost + TP?  (open question Q8)"]),
        mrs=("MRS", ["Railcar owner", "Sends under remittance", "Buys and receives the finished Pathfinder"]),
        bra=("Wabtec Digital Brazil", ["Buys components", "Importer of record in Brazil",
                                       "Coordinates refurbishment", "Sells the finished Pathfinder"]),
        l_pn="PN purchase + TP", l_rem="Railcar remittance\nno sale", l_sale="Sale of finished\nPathfinder",
        l_down="Finished Pathfinder\nIron Ore Car", l_up="Railcar +\ncomponents",
        pts_title="SEVEN DEAL POINTS",
        pts=["1  USA sells PNs + TP", "2  Digital Brazil imports", "3  MRS owns the railcar", "4  Remittance is not a sale",
             "5  Brazil sends railcar + components to Supermetal", "6  Supermetal returns finished Pathfinder + cost (TP?)",
             "7  Brazil sells the finished Pathfinder to MRS"],
        bands=[("IMPORT  ·  Q1–Q4", "II, import PIS/COFINS, IPI, ICMS-import and customs costs. Digital Brazil = importer of record", BLUE, BLUE_BG),
               ("REMITTANCE + REFURBISHMENT  ·  Q5–Q8", "Remitted railcar and applied materials kept separate. Validate CFOP, ICMS, IPI and ISS", ORANGE, ORANGE_BG),
               ("SALE TO MRS  ·  Q9–Q12", "Final price = imported components (PN + TP + import costs) + Supermetal cost + Digital margin + output taxes", GREEN, GREEN_BG)],
        foot="The remittance is not a sale of the railcar. Final tax treatment depends on entity, ownership, CFOP and tax classification — see the open tax questions on the next slide.",
        q_title="12 TAX QUESTIONS TO CLOSE BEFORE THE PATHFINDER IRON ORE CAR PRICE IS FINAL",
        q_sub="For the fiscal team to validate — none of these is answered in the current P&L",
        cols=[
            ("IMPORT  ·  Digital Brazil", BLUE, BLUE_BG, [
                ("Q1", "Transfer price: is the USA TP (13% / 17.2% is our placeholder) accepted as customs value between related parties? Which TP policy applies?"),
                ("Q2", "Ex-tarifário: can the Rack ship as one PN under one NCM covered by an ex (integrated systems can be refused)? Does the 1–2 year process fit the delivery date?"),
                ("Q3", "Software: are the licenses inside LXA / PTC / GoLinc part of the customs value, or separate payments (IRRF, CIDE, import PIS/COFINS on services)?"),
                ("Q4", "Credits: can Digital Brazil credit import ICMS and PIS/COFINS? Our P&L treats them as pass-through to MRS.")]),
            ("REMITTANCE + REFURBISHMENT", ORANGE, ORANGE_BG, [
                ("Q5", "Invoices: is the four-leg flow MRS → Digital → Supermetal → Digital valid as remittance for industrialization? Who is the ordering party, and are ICMS / IPI suspended on each leg?"),
                ("Q6", "Nature: is refurbishing a used car industrialization (ICMS / IPI) or a service (ISS)? Which municipality — Contagem 2% or another?"),
                ("Q7", "Value base: the railcar belongs to MRS. Is only the value added (labor + applied materials) taxed on return, and how is the car's value shown on the invoices?"),
                ("Q8", "Supermetal TP: does Supermetal charge cost only, or cost + margin / TP? Which invoice base applies?")]),
            ("SALE TO MRS", GREEN, GREEN_BG, [
                ("Q9", "What is sold: a new Pathfinder unit (goods: ICMS / IPI / PIS / COFINS) or a service on MRS's own car (ISS)? This sets every output tax."),
                ("Q10", "IPI: which NCM / TIPI rate applies to the finished car (wagon vs. control equipment), and who is the IPI taxpayer — Digital Brazil or Supermetal?"),
                ("Q11", "Client credits: can MRS credit ICMS (CIAP, 48 months) and PIS/COFINS? Does REIDI (PIS/COFINS suspension for rail infrastructure) apply? Drives MRS's ROI."),
                ("Q12", "Tax reform: how do CBS / IBS (2026–2033 transition) affect the price and the contract? Who bears tax changes after signing?")]),
        ],
        msg="Today's P&L assumes: II 17.96% blended (only 6 of 47 imported items on a confirmed NCM), PIS/COFINS 10.06%, ICMS 18%, IPI 9.69%, all passed through to MRS; Supermetal cost without ISS; no ex-tarifário.",
    ),
    "PT": dict(
        title="ESTRUTURA DO NEGÓCIO — PATHFINDER IRON ORE CAR",
        sub="Fluxo de propriedade, componentes, reforma e venda final",
        usa=("Wabtec Digital USA", ["Fornecedor dos PNs", "Vende componentes para a Digital Brasil"]),
        sup=("Supermetal", ["Recebe vagão + componentes", "Reforma e integra",
                            "Devolve o Pathfinder Iron Ore Car pronto", "Custo + TP?  (pergunta em aberto Q8)"]),
        mrs=("MRS", ["Proprietária do vagão", "Envia em remessa", "Compra e recebe o Pathfinder pronto"]),
        bra=("Wabtec Digital Brasil", ["Compra componentes", "Importadora oficial no Brasil",
                                       "Coordena a reforma", "Vende o Pathfinder pronto"]),
        l_pn="Compra dos PNs + TP", l_rem="Remessa do vagão\nsem venda", l_sale="Venda do Pathfinder\npronto",
        l_down="Pathfinder\nIron Ore Car pronto", l_up="Vagão +\ncomponentes",
        pts_title="SETE PONTOS DO NEGÓCIO",
        pts=["1  USA vende PNs + TP", "2  Digital Brasil importa", "3  MRS é dona do vagão", "4  Remessa não é venda",
             "5  Brasil envia vagão + componentes à Supermetal", "6  Supermetal devolve Pathfinder pronto + custo (TP?)",
             "7  Brasil vende o Pathfinder pronto à MRS"],
        bands=[("IMPORTAÇÃO  ·  Q1–Q4", "II, PIS/COFINS-import, IPI, ICMS-import e custos aduaneiros. Digital Brasil = importadora oficial", BLUE, BLUE_BG),
               ("REMESSA + REFORMA  ·  Q5–Q8", "Separar vagão remetido e materiais aplicados. Validar CFOP, ICMS, IPI e ISS", ORANGE, ORANGE_BG),
               ("VENDA À MRS  ·  Q9–Q12", "Preço final = componentes importados (PN + TP + custos de importação) + custo Supermetal + margem Digital + impostos de saída", GREEN, GREEN_BG)],
        foot="A remessa não é venda do vagão. O tratamento fiscal final depende da entidade, propriedade, CFOP e classificação tributária — ver as perguntas fiscais em aberto no próximo slide.",
        q_title="12 PERGUNTAS FISCAIS A FECHAR ANTES DE FIXAR O PREÇO DO PATHFINDER",
        q_sub="Para o time fiscal validar — nenhuma está respondida no P&L atual",
        cols=[
            ("IMPORTAÇÃO  ·  Digital Brasil", BLUE, BLUE_BG, [
                ("Q1", "Transfer price: o TP da USA (13% / 17,2% é nosso placeholder) é aceito como valor aduaneiro entre partes relacionadas? Qual política de TP vale?"),
                ("Q2", "Ex-tarifário: o Rack pode ir como um PN único em uma NCM coberta por ex (sistemas integrados podem ser recusados)? O processo de 1–2 anos cabe no prazo de entrega?"),
                ("Q3", "Software: as licenças dentro de LXA / PTC / GoLinc entram no valor aduaneiro ou são pagamentos separados (IRRF, CIDE, PIS/COFINS-import sobre serviços)?"),
                ("Q4", "Créditos: a Digital Brasil credita ICMS-import e PIS/COFINS-import? O nosso P&L trata como repasse à MRS.")]),
            ("REMESSA + REFORMA", ORANGE, ORANGE_BG, [
                ("Q5", "Notas: o fluxo de quatro pernas MRS → Digital → Supermetal → Digital vale como remessa para industrialização? Quem é o encomendante e ICMS / IPI ficam suspensos em cada perna?"),
                ("Q6", "Natureza: reformar um vagão usado é industrialização (ICMS / IPI) ou serviço (ISS)? Qual município — Contagem 2% ou outro?"),
                ("Q7", "Base de valor: o vagão é da MRS. Só o valor agregado (mão de obra + materiais aplicados) é tributado no retorno, e como o valor do vagão aparece nas notas?"),
                ("Q8", "TP da Supermetal: ela cobra só o custo, ou custo + margem / TP? Qual a base da nota?")]),
            ("VENDA À MRS", GREEN, GREEN_BG, [
                ("Q9", "O que se vende: uma unidade Pathfinder nova (mercadoria: ICMS / IPI / PIS / COFINS) ou um serviço no vagão da própria MRS (ISS)? Isso define todos os impostos de saída."),
                ("Q10", "IPI: qual NCM / alíquota TIPI vale para o vagão pronto (vagão vs. equipamento de controle) e quem é o contribuinte do IPI — Digital Brasil ou Supermetal?"),
                ("Q11", "Créditos do cliente: a MRS credita ICMS (CIAP, 48 meses) e PIS/COFINS? O REIDI (suspensão de PIS/COFINS em infraestrutura ferroviária) se aplica? Define o ROI da MRS."),
                ("Q12", "Reforma tributária: como CBS / IBS (transição 2026–2033) afetam o preço e o contrato? Quem arca com mudanças de imposto após a assinatura?")]),
        ],
        msg="O P&L de hoje assume: II 17,96% blended (só 6 de 47 itens importados com NCM confirmada), PIS/COFINS 10,06%, ICMS 18%, IPI 9,69%, tudo repassado à MRS; custo Supermetal sem ISS; sem ex-tarifário.",
    ),
}


def rgb(h):
    return RGBColor.from_string(h)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.MIDDLE, inset=(0.08, 0.05, 0.08, 0.05)):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(x) for x in inset]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.CENTER)
        if "after" in p:
            para.space_after = Pt(p["after"])
        for text, size, bold, color, font in p["runs"]:
            r = para.add_run()
            r.text = text
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
            r.font.color.rgb = rgb(color)


def run(text, size, bold=False, color=INK, font=FONT):
    return (text, size, bold, color, font)


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, paras, anchor=anchor, inset=(0, 0, 0, 0))
    return tb


def box(slide, x, y, w, h, fill, line, lw, paras, anchor=MSO_ANCHOR.MIDDLE, adj=0.10,
        inset=(0.08, 0.05, 0.08, 0.05), shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = adj
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    fill_tf(s.text_frame, paras, anchor=anchor, inset=inset)
    return s


def arrow(slide, x1, y1, x2, y2, color, width=2.4):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def entity(slide, x, y, w, h, spec, fill, line):
    name, lines = spec
    paras = [dict(runs=[run(name, 15, True, line, FONT_H)], after=3)]
    paras += [dict(runs=[run(t, 11)], align=PP_ALIGN.LEFT if False else PP_ALIGN.CENTER) for t in lines]
    return box(slide, x, y, w, h, fill, line, 1.6, paras, adj=0.12)


def chip(slide, cx, cy, w, h, text, color):
    lines = text.split("\n")
    paras = [dict(runs=[run(t, 10, True, color)]) for t in lines]
    return box(slide, cx - w / 2, cy - h / 2, w, h, "FFFFFF", color, 0.8, paras, adj=0.18, inset=(0.04, 0.02, 0.04, 0.02))


def diagram_slide(prs, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(s, M, 0.2, W - 2 * M, 0.5, [dict(runs=[run(t["title"], 22, True, INK)], align=PP_ALIGN.LEFT)])
    textbox(s, M + 0.02, 0.72, W - 2 * M, 0.3, [dict(runs=[run(t["sub"], 12, False, MUTED)], align=PP_ALIGN.LEFT)])

    sm_w, br_w, side_w = 3.5, 3.5, 2.75
    cx = W / 2
    sm_x, br_x = cx - sm_w / 2, cx - br_w / 2
    usa_x, mrs_x = M, W - M - side_w
    sm_y, sm_h = 1.08, 1.45
    row_y, row_h = 3.22, 1.45
    br_y, br_h = 3.15, 1.6

    # arrows first, chips and boxes on top
    arrow(s, usa_x + side_w, 3.95, br_x, 3.95, BLUE)
    arrow(s, mrs_x, 3.52, br_x + br_w, 3.52, GREEN)
    arrow(s, br_x + br_w, 4.35, mrs_x, 4.35, TEAL)
    xl, xr = cx - 0.6, cx + 0.6
    arrow(s, xl, sm_y + sm_h, xl, br_y, RED)
    arrow(s, xr, br_y, xr, sm_y + sm_h, ORANGE)

    entity(s, usa_x, row_y, side_w, row_h, t["usa"], BLUE_BG, BLUE)
    entity(s, sm_x, sm_y, sm_w, sm_h, t["sup"], ORANGE_BG, ORANGE)
    entity(s, mrs_x, row_y, side_w, row_h, t["mrs"], GREEN_BG, GREEN)
    entity(s, br_x, br_y, br_w, br_h, t["bra"], TEAL_BG, TEAL)

    gap_l = (usa_x + side_w + br_x) / 2
    gap_r = (br_x + br_w + mrs_x) / 2
    chip(s, gap_l, 3.95, 1.5, 0.5, t["l_pn"], BLUE)
    chip(s, gap_r, 3.52, 1.5, 0.55, t["l_rem"], GREEN)
    chip(s, gap_r, 4.35, 1.5, 0.55, t["l_sale"], TEAL)
    mid_y = (sm_y + sm_h + br_y) / 2
    chip(s, xl - 0.1 - 0.95, mid_y, 1.9, 0.52, t["l_down"], RED)
    chip(s, xr + 0.1 + 0.85, mid_y, 1.7, 0.52, t["l_up"], ORANGE)

    textbox(s, M, 4.9, W - 2 * M, 0.26, [dict(runs=[run(t["pts_title"], 11, True, MUTED)], align=PP_ALIGN.LEFT)])
    pt_colors = [BLUE, BLUE, GREEN, GREEN, TEAL, ORANGE, TEAL]
    n, gap = 7, 0.09
    pw = (W - 2 * M - (n - 1) * gap) / n
    for i, (txt, col) in enumerate(zip(t["pts"], pt_colors)):
        box(s, M + i * (pw + gap), 5.17, pw, 0.68, "FFFFFF", col, 1.1,
            [dict(runs=[run(txt, 10, True, INK)])], adj=0.15, inset=(0.05, 0.02, 0.05, 0.02))
    bw = (W - 2 * M - 2 * 0.2) / 3
    for i, (head, body, col, bg) in enumerate(t["bands"]):
        box(s, M + i * (bw + 0.2), 5.98, bw, 0.92, bg, col, 1.6,
            [dict(runs=[run(head, 11, True, col, FONT_H)], after=2), dict(runs=[run(body, 10)])],
            adj=0.12, inset=(0.1, 0.04, 0.1, 0.04))
    textbox(s, M, 6.98, W - 2 * M, 0.4, [dict(runs=[run(t["foot"], 10, True, RED)])], anchor=MSO_ANCHOR.TOP)
    return s


def questions_slide(prs, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(s, M, 0.2, W - 2 * M, 0.5, [dict(runs=[run(t["q_title"], 20, True, INK)], align=PP_ALIGN.LEFT)])
    textbox(s, M + 0.02, 0.72, W - 2 * M, 0.3, [dict(runs=[run(t["q_sub"], 12, False, MUTED)], align=PP_ALIGN.LEFT)])
    cw = (W - 2 * M - 2 * 0.2) / 3
    for i, (head, col, bg, qs) in enumerate(t["cols"]):
        paras = [dict(runs=[run(head, 14, True, col, FONT_H)], align=PP_ALIGN.LEFT, after=8)]
        for tag, q in qs:
            paras.append(dict(runs=[run(tag + "  ", 13, True, col), run(q, 13)], align=PP_ALIGN.LEFT, after=9))
        box(s, M + i * (cw + 0.2), 1.15, cw, 5.05, bg, col, 1.6, paras, anchor=MSO_ANCHOR.TOP, adj=0.04,
            inset=(0.15, 0.12, 0.15, 0.08))
    box(s, M, 6.38, W - 2 * M, 0.85, INK, None, 0, [dict(runs=[run(t["msg"], 11.5, True, "FFFFFF")], align=PP_ALIGN.LEFT)],
        adj=0.08, inset=(0.2, 0.06, 0.2, 0.06))
    return s


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
for lang in ("EN", "PT"):
    diagram_slide(prs, TXT[lang])
    questions_slide(prs, TXT[lang])
prs.save(OUT)
print("saved", OUT)
