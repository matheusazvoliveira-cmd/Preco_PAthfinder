"""Builds MRS_Pathfinder_Simulation_Local_Taxes_Bilingual.pptx (slide 1 EN, slide 2 PT).
Scenario numbers come from the Simulation workbook; the US landed reference uses
the cached all-taxes total and FX from the complete import-cost workbook.
Run: PYTHONIOENCODING=utf-8 python build_simulation_slide.py"""
from pathlib import Path
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from lxml import etree

OUT = Path(__file__).parent / "MRS_Pathfinder_Simulation_Local_Taxes_Bilingual.pptx"
RED, DARK_BLUE, BLUE, INK, MUTED = "D70010", "003158", "3685C0", "2B2A2F", "757588"
AQUA, YELLOW = "00847E", "FFAB18"
PANEL, LIGHT, WHITE = "F8F8F8", "E1E0D9", "FFFFFF"
FONT, FONT_H = "Aptos", "Aptos Display"
NET = [945, 845, 908, 813]; TAX = [461, 412, 442, 396]; FINAL = [1406, 1257, 1350, 1210]
source = openpyxl.load_workbook(
    Path(__file__).parent / "Planilha de Custos de Importacao - Pathfinder completo- Set 26.xlsx",
    read_only=True, data_only=True)["Pathfinder Pronto"]
assert source["H24"].value == 900000, "US reference sticker price changed"
assert source["I56"].value and source["I22"].value, "Missing cached landed total or FX"
FX = source["I22"].value


def usd_k(brl):
    return brl / FX / 1000


# I56 = goods + international freight + insurance + customs expenses + II/IPI/PIS/COFINS/ICMS.
REF_PROD = usd_k(source["I26"].value)          # product / sticker price
REF_FRT = usd_k(source["I27"].value + source["I28"].value)  # international freight + insurance
REF_EXP = usd_k(source["I35"].value)           # customs / import expenses
REF_TAX = usd_k(source["I56"].value - source["I26"].value - source["I27"].value
                - source["I28"].value - source["I35"].value)
REF = usd_k(source["I56"].value)               # USD thousands, like the other columns
assert abs(REF_PROD - 900) < 1e-6, "US reference product price is no longer US$900k"
assert abs(REF_PROD + REF_FRT + REF_EXP + REF_TAX - REF) < 1e-6, "Segments do not sum to the landed total"
YMAX = 1900
# inner plot area as fractions of chart frame
PX, PY, PW, PH = 0.03, 0.10, 0.76, 0.66
CX, CY, CW, CH = 0.45, 1.25, 7.7, 4.0


def k_en(v):
    return "US$" + format(v, ",") + "k"


def k_pt(v):
    return "US$" + format(v, ",").replace(",", ".") + "k"


def n_pt(v, d=1):
    return format(v, ",.%df" % d).replace(",", "X").replace(".", ",").replace("X", ".")


T = {
    "EN": dict(
        title="Even with zero import duty and one battery set, local taxes keep the Pathfinder expensive",
         cats=["1  Base\n(2 battery sets,\nimport duty in)", "2  No import\nduty (II = 0)", "3  One battery\nset (II in)", "4  No II +\none battery set", "US reference\n(US$900k US price\nlanded in Brazil)"],
        s_net="Our price (net, before local taxes)", s_tax="Local taxes (PIS/COFINS + ICMS + IPI)",
         s_ref=["US reference: US price (product)", "US reference: freight + insurance",
                "US reference: customs / import expenses", "US reference: all import taxes"], fmt=k_en,
         call_h="How the US$900k becomes " + k_en(round(REF)),
         call_rows=[("US price (product, unchanged)", REF_PROD, MUTED),
                    ("+ Freight + insurance", REF_FRT, AQUA),
                    ("+ Customs / import expenses", REF_EXP, YELLOW),
                    ("+ All import taxes (II, IPI, PIS/COFINS, ICMS)", REF_TAX, BLUE)],
         call_tot="= Landed in Brazil",
        tiles=[("US$149k", "Effect of removing ALL import duty (II) on the final client price"),
               ("US$461k", "Local taxes on top of our net price in the base case = 3.1x the II effect"),
             (k_en(round(REF)), "US$900k US price + US$%dk freight/insurance + US$%dk import expenses + US$%dk import taxes"
              % (round(REF_FRT), round(REF_EXP), round(REF_TAX)))],
        msg="Removing all import duty and using one battery set still leaves a final price of US$1,210k. What makes the Pathfinder expensive for MRS is the local Brazilian taxes, not the import duty.",
        note="II matters more than its face value because it sits inside the cost: it is multiplied by the margin (1/(1-45%)) and by the local tax cascade (~2.7x), and local taxes still outweigh it.\n"
             "Assumptions (scenarios): TDM 45%; FX 5.10 BRL/USD; II 17.96% blended (only 6 of 47 imported items have a confirmed NCM, the rest at a worst-case 18%); PIS/COFINS 10.06%, ICMS 18%, IPI 9.69% passed through to MRS; Supermetal refurbishment included, no ISS; 27 items still without a quote are excluded. US landed reference: complete import-cost workbook ('Pathfinder Pronto' sheet), FX 5.15 BRL/USD, product US$"
             + format(REF_PROD, ",.0f") + "k + freight/insurance US$" + format(REF_FRT, ",.1f")
             + "k + import expenses US$" + format(REF_EXP, ",.1f") + "k + all import taxes US$"
             + format(REF_TAX, ",.1f") + "k = US$" + format(REF, ",.1f")
             + "k; not the same scope as the scenario estimates."),
    "PT": dict(
         title="Mesmo sem II, os impostos locais encarecem o Pathfinder",
         cats=["1  Base\n(2 jogos de baterias,\nII incluído)", "2  Sem imposto de\nimportação (II = 0)", "3  Um jogo de\nbaterias (II incluído)", "4  Sem II +\num jogo de baterias", "Referência EUA\n(preço EUA US$900k\nposto no Brasil)"],
        s_net="Nosso preço (líquido, antes dos impostos locais)", s_tax="Impostos locais (PIS/COFINS + ICMS + IPI)",
         s_ref=["Referência EUA: preço EUA (produto)", "Referência EUA: frete + seguro",
                "Referência EUA: despesas aduaneiras", "Referência EUA: todos os impostos de importação"], fmt=k_pt,
         call_h="Como os US$900k viram " + k_pt(round(REF)),
         call_rows=[("Preço EUA (produto, inalterado)", REF_PROD, MUTED),
                    ("+ Frete + seguro", REF_FRT, AQUA),
                    ("+ Despesas aduaneiras / importação", REF_EXP, YELLOW),
                    ("+ Todos os impostos de importação (II, IPI, PIS/COFINS, ICMS)", REF_TAX, BLUE)],
         call_tot="= Posto no Brasil",
        tiles=[("US$149k", "Efeito de zerar TODO o imposto de importação (II) no preço final ao cliente"),
               ("US$461k", "Impostos locais sobre o nosso preço líquido no caso base = 3,1x o efeito do II"),
             (k_pt(round(REF)), "Preço EUA US$900k + US$%dk frete/seguro + US$%dk despesas + US$%dk impostos de importação"
              % (round(REF_FRT), round(REF_EXP), round(REF_TAX)))],
        msg="Zerar todo o imposto de importação e usar um só jogo de baterias ainda deixa o preço final em US$1.210k. O que encarece o Pathfinder para a MRS são os impostos locais brasileiros, não o imposto de importação.",
        note="O II pesa mais que o seu valor nominal porque está dentro do custo: é multiplicado pela margem (1/(1-45%)) e pela cascata de impostos locais (~2,7x), e mesmo assim os impostos locais pesam mais.\n"
             "Premissas (cenários): TDM 45%; câmbio 5,10 BRL/USD; II 17,96% médio (só 6 de 47 itens importados têm NCM confirmada, o restante a 18% pior caso); PIS/COFINS 10,06%, ICMS 18%, IPI 9,69% repassados à MRS; reforma da Supermetal incluída, sem ISS; 27 itens ainda sem cotação excluídos. Referência EUA posta: planilha completa de importação (aba 'Pathfinder Pronto'), câmbio 5,15 BRL/USD, produto US$"
             + n_pt(REF_PROD, 0) + "k + frete/seguro US$" + n_pt(REF_FRT) + "k + despesas de importação US$"
             + n_pt(REF_EXP) + "k + todos os impostos de importação US$" + n_pt(REF_TAX) + "k = US$" + n_pt(REF)
             + "k; escopo diferente das simulações."),
}


def rgb(h):
    return RGBColor.from_string(h)


def fill(tf, paras, anchor=MSO_ANCHOR.MIDDLE, inset=(0.1, 0.05, 0.1, 0.05)):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(x) for x in inset]
    for i, (runs, align) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, color, font in runs:
            r = p.add_run()
            r.text = text
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
            r.font.color.rgb = rgb(color)


def box(sl, x, y, w, h, fillc, paras, line=None, anchor=MSO_ANCHOR.MIDDLE, inset=(0.1, 0.05, 0.1, 0.05)):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fillc:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fillc)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    fill(s.text_frame, paras, anchor, inset)
    return s


def bar(sl, x, y, w, h, col):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(col)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def build(prs, lang):
    t = T[lang]
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bar(sl, 0.45, 0.3, 0.09, 0.8, RED)
    box(sl, 0.6, 0.3, 12.3, 0.8, None, [([(t["title"], 26, True, INK, FONT_H)], PP_ALIGN.LEFT)])
    cd = CategoryChartData()
    cd.categories = t["cats"]
    cd.add_series(t["s_net"], NET + [None])
    cd.add_series(t["s_tax"], TAX + [None])
    for name, v in zip(t["s_ref"], (REF_PROD, REF_FRT, REF_EXP, REF_TAX)):
        cd.add_series(name, [None] * 4 + [v])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(CX), Inches(CY), Inches(CW), Inches(CH), cd)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(11)
    ch.font.color.rgb = rgb(INK)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    leg = ch.legend._element
    ns_c = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    pos = leg.find("{%s}legendPos" % ns_c)
    at = list(leg).index(pos) + 1 if pos is not None else 0
    for idx in range(2, 6):  # the four US-reference segments are explained by the callout, not the legend
        leg.insert(at, etree.fromstring(
            '<c:legendEntry xmlns:c="%s"><c:idx val="%d"/><c:delete val="1"/></c:legendEntry>' % (ns_c, idx)))
        at += 1
    va = ch.value_axis
    va.minimum_scale = 0
    va.maximum_scale = YMAX
    va.visible = False
    va.has_major_gridlines = False
    ch.category_axis.tick_labels.font.size = Pt(10)
    ch.category_axis.format.line.color.rgb = rgb(MUTED)
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.overlap = 100
    for i, (ser, col) in enumerate(zip(plot.series, (RED, BLUE, MUTED, AQUA, YELLOW, BLUE))):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)
        if i in (3, 4):
            continue  # freight and import expenses are too thin to label inside; the callout carries them
        dl = ser.data_labels
        dl.show_value = True
        dl.number_format = '"US$"0"k"'
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.CENTER
        dl.font.size = Pt(14 if i < 2 else 12)
        dl.font.bold = True
        dl.font.color.rgb = rgb(WHITE)
    pa = ch._chartSpace.chart.plotArea
    ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    old = pa.find("{%s}layout" % ns)
    if old is not None:
        pa.remove(old)
    lay = etree.fromstring(
        '<c:layout xmlns:c="%s"><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
        '<c:x val="%s"/><c:y val="%s"/><c:w val="%s"/><c:h val="%s"/></c:manualLayout></c:layout>' % (ns, PX, PY, PW, PH))
    pa.insert(0, lay)
    px0 = CX + PX * CW
    pw = PW * CW
    py0 = CY + PY * CH
    ph = PH * CH

    def yv(v):
        return py0 + ph * (1 - v / YMAX)

    for i, f in enumerate(FINAL + [round(REF)]):
        cx = px0 + pw * (i + 0.5) / 5
        box(sl, cx - 0.55, yv(f) - 0.5, 1.1, 0.46, None,
            [([(t["fmt"](f), 14, True, INK, FONT_H)], PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.BOTTOM, inset=(0, 0, 0, 0.02))
    tx, tw, th, gap = 8.45, 4.43, 1.18, 0.13
    # Composition callout for the US reference column, to the right of the plot area.
    xr = px0 + pw * 4.5 / 5 + (pw / 5) / (1 + plot.gap_width / 100.0) / 2
    cox = xr + 0.29
    cow = tx - 0.05 - cox
    ymid = (yv(REF_PROD) + yv(REF_PROD + REF_FRT + REF_EXP)) / 2
    bar(sl, xr, ymid - 0.01, cox - xr, 0.02, AQUA)
    coy = 1.30
    box(sl, cox, coy, cow, 0.34, DARK_BLUE, [([(t["call_h"], 11, True, WHITE, FONT)], PP_ALIGN.LEFT)],
        inset=(0.1, 0.03, 0.08, 0.03))
    coy += 0.36
    for (label, val, col), h in zip(t["call_rows"], (0.42, 0.42, 0.46, 0.66)):
        bar(sl, cox, coy, 0.09, h, col)
        box(sl, cox + 0.09, coy, cow - 0.09, h, PANEL,
            [([(t["fmt"](round(val)), 12, True, INK, FONT_H)], PP_ALIGN.LEFT),
             ([(label, 8, False, INK, FONT)], PP_ALIGN.LEFT)],
            anchor=MSO_ANCHOR.TOP, inset=(0.08, 0.03, 0.06, 0.02))
        coy += h + 0.02
    box(sl, cox, coy, cow, 0.40, None,
        [([(t["fmt"](round(REF)) + "  ", 12, True, INK, FONT_H), (t["call_tot"], 8, False, MUTED, FONT)], PP_ALIGN.LEFT)],
        line=LIGHT, inset=(0.08, 0.02, 0.06, 0.02))
    for i, (big, small) in enumerate(t["tiles"]):
        y = 1.25 + i * (th + gap)
        bar(sl, tx, y, 0.08, th, RED if i == 1 else BLUE)
        box(sl, tx + 0.08, y, tw - 0.08, th, PANEL,
            [([(big, 30, True, INK, FONT_H)], PP_ALIGN.LEFT), ([(small, 12, False, INK, FONT)], PP_ALIGN.LEFT)],
            inset=(0.18, 0.04, 0.12, 0.04))
    box(sl, 0.45, 5.4, 12.43, 0.85, DARK_BLUE, [([(t["msg"], 16, True, WHITE, FONT)], PP_ALIGN.LEFT)], inset=(0.25, 0.05, 0.25, 0.05))
    a, b = t["note"].split("\n")
    box(sl, 0.45, 6.32, 12.43, 1.0, None,
        [([(a, 10, False, INK, FONT)], PP_ALIGN.LEFT), ([(b, 10, False, MUTED, FONT)], PP_ALIGN.LEFT)],
        anchor=MSO_ANCHOR.TOP, inset=(0, 0, 0, 0))
    return sl


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
build(prs, "EN")
build(prs, "PT")
prs.save(OUT)
print("saved", OUT)
